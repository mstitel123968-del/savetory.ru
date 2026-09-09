"""Local exports. Write atomically, keep custom fields and neutralize formulas."""
import os
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

STATUS = {'keep': 'Храню', 'sell': 'Готов продать', 'exchange': 'Готов обменять', 'search': 'Ищу такой же', 'sold': 'Продано'}


def fields_for(store, card):
    rubric = store.rubric(card.get('rubric_id')) or {}
    values = card.get('values', {})
    known = set()
    for field in rubric.get('fields', []):
        if field.get('type') == 'image':
            continue
        key = field['id']
        known.add(key)
        yield key, field['label'], str(values.get(key, ''))
    for key, value in values.items():
        if key not in known:
            yield key, {'title': 'Наименование', 'year': 'Год'}.get(key, key), str(value)


def export_collection(store, cards, destination, kind):
    destination = Path(destination)
    fd, staging = tempfile.mkstemp(prefix='.sklad-export-', suffix='.'+kind, dir=destination.parent)
    os.close(fd)
    try:
        if kind == 'xlsx':
            write_excel(store, cards, staging)
        elif kind == 'pdf':
            write_pdf(store, cards, staging)
        else:
            raise ValueError('Неизвестный формат экспорта')
        os.replace(staging, destination)
    finally:
        if os.path.exists(staging):
            os.unlink(staging)


def write_excel(store, cards, destination):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    columns = []
    for card in cards:
        for key, label, _ in fields_for(store, card):
            if (key, label) not in columns:
                columns.append((key, label))
    book = Workbook()
    sheet = book.active
    sheet.title = 'Коллекция'
    headers = ['Рубрика'] + [label for _, label in columns] + ['Изображения']
    def literal(row, col, value):
        text = ILLEGAL_CHARACTERS_RE.sub('', str(value))
        if len(text) > 32767:
            raise ValueError('Поле превышает ограничение Excel в 32767 символов. Выберите PDF.')
        cell = sheet.cell(row, col, text)
        cell.data_type = 's'  # Titles starting with =, +, - or @ are never formulas.
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        return cell
    for col, title in enumerate(headers, 1):
        cell = literal(1, col, title)
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
        sheet.column_dimensions[get_column_letter(col)].width = 28
    for row, card in enumerate(cards, 2):
        rubric = store.rubric(card.get('rubric_id')) or {}
        mapped = {(key, label): value for key, label, value in fields_for(store, card)}
        images = '\n'.join(str(store.resolve_image(path) or path) for path in card.get('images', []))
        values = [rubric.get('name', 'Без рубрики')]
        values += [mapped.get(column, '') for column in columns] + [images]
        for col, value in enumerate(values, 1):
            cell = literal(row, col, value)
            if row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EDF3F8')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    book.save(destination)


def write_pdf(store, cards, destination):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
    fonts = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    font = next((fonts / name for name in ('segoeui.ttf', 'arial.ttf') if (fonts / name).exists()), None)
    if font is None:
        raise ValueError('Не найден системный шрифт с поддержкой кириллицы для PDF.')
    pdfmetrics.registerFont(TTFont('SKladText', str(font)))
    body = ParagraphStyle('body', fontName='SKladText', fontSize=10, leading=15, textColor=colors.HexColor('#243B53'), splitLongWords=True)
    title = ParagraphStyle('title', parent=body, fontSize=20, leading=26, spaceAfter=12)
    label_style = ParagraphStyle('label', parent=body, textColor=colors.HexColor('#65758B'), spaceBefore=9, spaceAfter=3)
    story = []
    def paragraph(text, style=body):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), style)
    if not cards:
        story = [paragraph('Коллекция пуста', title)]
    for index, card in enumerate(cards):
        if index:
            story.append(PageBreak())
        rubric = store.rubric(card.get('rubric_id')) or {}
        story.append(paragraph(card.get('values', {}).get('title', 'Без названия'), title))
        story.append(paragraph(rubric.get('name', 'Без рубрики')))
        story.append(Spacer(1, 14))
        for source in card.get('images', [])[:5]:
            path = store.resolve_image(source)
            if path and path.exists():
                try:
                    from PIL import Image as PILImage
                    with PILImage.open(path) as probe:
                        probe.verify()
                    photo = Image(str(path))
                    ratio = min(460/photo.imageWidth, 200/photo.imageHeight)
                    photo.drawWidth, photo.drawHeight = photo.imageWidth*ratio, photo.imageHeight*ratio
                    photo.hAlign = 'LEFT'
                    story.extend([photo, Spacer(1, 10)])
                except (OSError, ValueError):
                    story.append(paragraph('Изображение недоступно'))
        for key, label, value in fields_for(store, card):
            if key != 'title':
                story.extend([paragraph(label, label_style), paragraph(value or '—')])
    doc = SimpleDocTemplate(str(destination), pagesize=A4, leftMargin=42, rightMargin=42, topMargin=42, bottomMargin=42, title='СКлад — коллекция', author='СКлад')
    def footer(canvas, _doc):
        canvas.setFont('SKladText', 8)
        canvas.setFillColor(colors.HexColor('#65758B'))
        canvas.drawString(42, 24, 'СКлад · Экспорт коллекции')
        canvas.drawRightString(A4[0]-42, 24, str(canvas.getPageNumber()))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
