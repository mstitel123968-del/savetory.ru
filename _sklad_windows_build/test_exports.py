import tempfile
import unittest
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
from storage import Storage
from exports import export_collection


class ExportTests(unittest.TestCase):
    def test_roundtrip_and_atomic_failure(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            store = Storage(root, root, root / 'data')
            image = root / 'photo.png'
            Image.new('RGB', (240, 120), '#336699').save(image)
            card = {'id': 'test', 'rubric_id': 'coins', 'status': 'keep', 'images': [str(image)],
                    'values': {'title': '=1+1', 'description': 'Описание <предмета> & заметки\n' * 160,
                               'custom_rare': 'Редкое значение: Ёжик'}}
            xlsx, pdf = root/'out.xlsx', root/'out.pdf'
            export_collection(store, [card], xlsx, 'xlsx')
            book = load_workbook(xlsx)
            cells = list(book.active.iter_rows())[1]
            title = next(c for c in cells if c.value == '=1+1')
            self.assertEqual(title.data_type, 's')
            self.assertIn('Редкое значение: Ёжик', [c.value for c in cells])
            export_collection(store, [card], pdf, 'pdf')
            from pypdf import PdfReader
            document = PdfReader(pdf)
            text = '\n'.join(page.extract_text() for page in document.pages)
            self.assertIn('Ёжик', text)
            self.assertIn('Описание <предмета> & заметки', text)
            self.assertGreater(len(document.pages), 1)
            before = xlsx.read_bytes()
            with self.assertRaises(ValueError):
                export_collection(store, [card], xlsx, 'bad-format')
            self.assertEqual(xlsx.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
